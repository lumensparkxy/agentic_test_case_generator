from copy import deepcopy
import csv
from datetime import datetime, timezone
import io
import json
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.main import app, get_current_user
from app.models import AuthUser
from app.contracts.projects import QaProjectDetail
from app.services.workflow_project_service import ProjectPermissionError


def _test_case():
    return {
        "id": "TC-001",
        "title": '=HYPERLINK("https://example.test")',
        "status": "Ready",
        "steps": [{"step": 1, "action": "Open checkout", "expected": "Checkout is visible"}],
        "linked_requirement_ids": ["REQ-001"],
        "scenario_refs": ["SCN-1"],
    }


def _project(approved=False):
    now = datetime.now(timezone.utc)

    def snapshot(stage, identity, payload, **kwargs):
        return {
            "snapshot_id": identity,
            "project_id": "project-1",
            "stage": stage,
            "version": 2,
            "project_revision": 9,
            "operation": "fixture",
            "created_at": now,
            "payload": payload,
            **kwargs,
        }

    return QaProjectDetail(
        project_id="project-1",
        name="Export QA",
        owner_user_id="export-user",
        current_revision=9,
        created_at=now,
        updated_at=now,
        stage_state={"test_cases": {"current_snapshot_id": "tests-2", "approved": approved, "version": 2}},
        current_snapshots={
            "requirements": snapshot(
                "requirements",
                "req-2",
                {
                    "requirements": [
                        {
                            "id": "REQ-001",
                            "text": "Visible checkout",
                            "original_requirement_ids": ["ROOM-101"],
                            "sources": [
                                {
                                    "source_id": "source-1",
                                    "source_version": "hash-1",
                                    "import_id": "import-1",
                                    "label": "Room",
                                    "excerpt": "+SUM(1,2) Zürich",
                                    "excerpt_verified": True,
                                    "original_requirement_ids": ["ROOM-101"],
                                }
                            ],
                        },
                        {"id": "REQ-002", "text": "Cancel checkout", "sources": []},
                    ]
                },
            ),
            "test_cases": snapshot(
                "test_cases",
                "tests-2",
                {
                    "test_cases": [_test_case()],
                    "review": {"approved": approved, "score": 100 if approved else 55, "threshold": 90},
                    "generation_tasks": [] if approved else [{"reason": "Cancellation missing"}],
                    "coverage_plan": [{"requirement_id": "REQ-001", "scenarios": [{"id": "SCN-1"}]}]
                    if approved
                    else [{"requirement_id": "REQ-001", "scenarios": [{"id": "SCN-1"}, {"id": "SCN-2"}]}],
                },
                approved=approved,
                metadata={"source_requirements_snapshot_id": "req-2"},
            ),
        },
    )


def _payload(**kwargs):
    return {
        "project_id": "project-1",
        "base_project_revision": 9,
        "source_snapshot_id": "tests-2",
        "test_cases": [_test_case()],
        "approved": True,
        "review": {"approved": True},
        "draft_override_requested": True,
        "draft_override_reason": "@review Zürich",
        **kwargs,
    }


class ExportEndpointTests(unittest.TestCase):
    def setUp(self):
        app.dependency_overrides[get_current_user] = lambda: AuthUser(sub="export-user", name="Export QA")
        self.addCleanup(app.dependency_overrides.clear)
        for target, result in (("start_workflow_run", "workflow-1"), ("complete_workflow_run", None), ("record_usage_event", "event-1")):
            patcher = patch(f"app.routers.export.{target}", return_value=result)
            patcher.start()
            self.addCleanup(patcher.stop)

    def export(self, format="json", payload=None, project=None):
        with (
            patch("app.routers.export.get_project", return_value=project or _project()) as get,
            patch("app.routers.export.append_stage_snapshot") as append,
            TestClient(app) as client,
        ):
            result = client.post(f"/export/{format}", json=payload or _payload())
        return result, append, get

    def test_all_formats_preserve_draft_evidence_and_case_identity(self):
        for format in ("json", "csv", "excel"):
            with self.subTest(format=format):
                response, append, get = self.export(format)
                self.assertEqual(response.status_code, 200, response.text[:300] if format != "excel" else response.status_code)
                metadata = append.call_args.kwargs["payload"]["audit_metadata"]
                self.assertTrue(metadata["is_draft"])
                self.assertEqual(metadata["machine_review_status"], "failed")
                self.assertEqual(metadata["human_review_status"], "unavailable")
                self.assertFalse(metadata["suite_complete"])
                self.assertEqual(metadata["unresolved_task_count"], 1)
                self.assertEqual(metadata["coverage"]["requirements"]["missing_ids"], ["REQ-002"])
                self.assertEqual(metadata["coverage"]["scenarios"]["missing_ids"], ["SCN-2"])
                self.assertIsNone(metadata["coverage"]["behavioral_assessment"])
                self.assertEqual(metadata["source_mappings"][0]["original_requirement_ids"], ["ROOM-101"])
                self.assertEqual(metadata["execution_status"], "not_executed")
                self.assertEqual(metadata["snapshot_id"], "tests-2")
                self.assertEqual(metadata["draft_override_reason"], "@review Zürich")
                self.assertFalse(append.call_args.kwargs["approved"])
                self.assertEqual(append.call_args.kwargs["source_snapshot_id"], "tests-2")
                get.assert_called_once()
                if format == "json":
                    body = response.json()
                    self.assertEqual(body["export_format"], "test_cases_v1")
                    self.assertEqual(body["total_count"], 1)
                    self.assertEqual(body["test_cases"][0]["title"], _test_case()["title"])
                    self.assertEqual(body["test_cases"][0]["status"], "Ready")
                    self.assertTrue(body["audit_metadata"]["is_draft"])
                elif format == "csv":
                    rows = list(csv.DictReader(io.StringIO(response.text)))
                    self.assertEqual(len(rows), 1)
                    self.assertEqual(rows[0]["ID"], "TC-001")
                    self.assertTrue(rows[0]["Title"].startswith("'="))
                    self.assertEqual(rows[0]["Audit draft_override_reason"], "'@review Zürich")
                    self.assertEqual(json.loads(rows[0]["Audit coverage"])["scenarios"]["missing_ids"], ["SCN-2"])
                else:
                    wb = load_workbook(io.BytesIO(response.content))
                    self.assertEqual(wb.sheetnames, ["Test Cases", "Audit", "Sources", "Coverage"])
                    self.assertEqual(wb["Test Cases"]["A2"].value, "TC-001")
                    self.assertEqual(wb["Test Cases"]["F2"].value, "Ready")
                    self.assertNotEqual(wb["Test Cases"]["B2"].data_type, "f")
                    audit = {row[0].value: row[1].value for row in wb["Audit"].iter_rows(min_row=2)}
                    self.assertEqual(audit["machine_review_status"], "failed")
                    self.assertEqual(audit["draft_override_reason"], "'@review Zürich")
                    self.assertEqual(wb["Sources"]["F2"].value, "'+SUM(1,2) Zürich")
                    self.assertFalse(any(cell.data_type == "f" for sheet in wb for row in sheet for cell in row))

    def test_approved_export_uses_matching_saved_review_not_client_claim(self):
        project = _project(True)
        project.current_snapshots["requirements"].payload["requirements"].pop()
        payload = _payload(approved=False, review={"approved": False}, draft_override_requested=False, draft_override_reason=None)
        response, append, _ = self.export(payload=payload, project=project)
        self.assertEqual(response.status_code, 200, response.text)
        metadata = response.json()["audit_metadata"]
        self.assertFalse(metadata["is_draft"])
        self.assertEqual(metadata["machine_review_status"], "approved")
        self.assertEqual(metadata["snapshot_version"], 2)
        self.assertTrue(append.call_args.kwargs["approved"])
        self.assertEqual(metadata["human_review_status"], "unavailable")

    def test_forged_approval_empty_reason_and_unknown_review_cannot_bypass_draft_gate(self):
        for variant in ("override_off", "empty_reason", "unknown_review", "incomplete_approved"):
            project = _project(variant == "incomplete_approved")
            payload = _payload(draft_override_requested=False)
            if variant == "empty_reason":
                payload.update(draft_override_requested=True, draft_override_reason=" \n ")
            if variant == "unknown_review":
                del project.current_snapshots["test_cases"].payload["review"]
            response, append, _ = self.export(payload=payload, project=project)
            self.assertEqual(response.status_code, 422, (variant, response.text))
            append.assert_not_called()

    def test_stale_mismatched_and_forged_cases_conflict_before_export(self):
        for variant in ("revision", "missing_revision", "snapshot", "case", "stale", "old_source", "unknown_snapshot"):
            project = _project()
            payload = _payload()
            if variant == "revision":
                payload["base_project_revision"] = 8
            if variant == "missing_revision":
                payload.pop("base_project_revision")
            if variant == "snapshot":
                payload["source_snapshot_id"] = "old"
            if variant == "case":
                payload["test_cases"][0]["title"] = "Forged"
            if variant == "stale":
                project.stage_state["test_cases"].stale = True
            if variant == "old_source":
                project.current_snapshots["test_cases"].metadata["source_requirements_snapshot_id"] = "old"
            if variant == "unknown_snapshot":
                project.current_snapshots.pop("test_cases")
            response, append, _ = self.export(payload=payload, project=project)
            self.assertEqual(response.status_code, 409, (variant, response.text))
            append.assert_not_called()

    def test_unbound_draft_and_legacy_provenance_are_explicitly_unknown(self):
        payload = _payload(project_id=None, source_snapshot_id=None)
        response, append, get = self.export(payload=payload)
        self.assertEqual(response.status_code, 200, response.text)
        metadata = response.json()["audit_metadata"]
        self.assertIsNone(metadata["snapshot_id"])
        self.assertIsNone(metadata["suite_complete"])
        self.assertEqual(metadata["machine_review_status"], "unavailable")
        self.assertEqual(metadata["source_mappings"], [])
        self.assertEqual(metadata["execution_status"], "unknown")
        get.assert_not_called()
        append.assert_not_called()
        project = _project()
        project.current_snapshots["test_cases"].metadata.clear()
        response, _, _ = self.export(project=project)
        self.assertEqual(response.json()["audit_metadata"]["source_mappings"], [])

    def test_only_matching_execution_runs_are_evidence(self):
        project = _project()
        now = datetime.now(timezone.utc)
        from app.contracts.projects import QaProjectExecutionRun

        project.execution_runs = [
            QaProjectExecutionRun(
                run_record_id="r",
                project_id="project-1",
                run_id="past-run",
                target_environment="qa",
                project_revision=7,
                status="passed",
                source_snapshot_id="old-tests",
                created_at=now,
            )
        ]
        response, _, _ = self.export(project=project)
        self.assertEqual(response.json()["audit_metadata"]["execution_run_ids"], [])
        self.assertEqual(response.json()["audit_metadata"]["execution_status"], "unknown")
        project.execution_runs[0].source_snapshot_id = "tests-2"
        project.execution_runs[0].status = "failed"
        response, append, _ = self.export(project=project)
        self.assertEqual(response.json()["audit_metadata"]["execution_status"], "failed")
        self.assertEqual(append.call_args.kwargs["payload"]["evidence"]["execution_run_ids"], ["past-run"])

    def test_unauthorized_project_cannot_export(self):
        with (
            patch("app.routers.export.get_project", side_effect=ProjectPermissionError("project-1")),
            patch("app.routers.export.export_to_json") as export,
            TestClient(app) as client,
        ):
            response = client.post("/export/json", json=_payload())
        self.assertEqual(response.status_code, 403)
        export.assert_not_called()

    def test_human_review_is_independent_and_bound_to_the_same_snapshot(self):
        project = _project(True)
        project.current_snapshots["requirements"].payload["requirements"].pop()
        project.stage_state["test_cases"].metadata["latest_human_review"] = {"snapshot_id": "tests-2", "decision": "request_changes"}
        response, _, _ = self.export(project=project)
        self.assertTrue(response.json()["audit_metadata"]["is_draft"])
        self.assertEqual(response.json()["audit_metadata"]["human_review_status"], "changes_requested")
        self.assertEqual(response.json()["audit_metadata"]["machine_review_status"], "approved")
        project.stage_state["test_cases"].metadata["latest_human_review"]["snapshot_id"] = "older"
        response, _, _ = self.export(project=project)
        self.assertEqual(response.json()["audit_metadata"]["human_review_status"], "unavailable")
        project.stage_state["test_cases"].metadata["latest_human_review"] = {"snapshot_id": "tests-2", "decision": "approve"}
        response, _, _ = self.export(project=project)
        self.assertEqual(response.json()["audit_metadata"]["human_review_status"], "approved")
        response, append, _ = self.export(payload=_payload(project_id=None, source_snapshot_id=None, draft_override_requested=False))
        self.assertEqual(response.status_code, 422)
        append.assert_not_called()

    def test_placeholder_remains_blocked_with_override(self):
        payload = _payload(project_id=None, source_snapshot_id=None)
        payload["test_cases"][0]["steps"] = []
        response, append, _ = self.export(payload=payload)
        self.assertEqual(response.status_code, 422, response.text)
        append.assert_not_called()
